import html
import logging
import mimetypes
import os
import tempfile

import pandas as pd
from docx import Document as DocxDocument
from fastapi import HTTPException
from fastapi.responses import FileResponse, HTMLResponse, StreamingResponse
from starlette.responses import Response

from apps.shared.db.models.knowledge import Document
from apps.shared.services.egress_guard import (
    EgressGuardError,
    safe_http_request,
    safe_quote_filename,
)
from apps.shared.services.outbound_operation_policy import KNOWLEDGE_DOCUMENT_FETCH

logger = logging.getLogger(__name__)

SAFE_INLINE_FILE_EXTENSIONS = {".md", ".pdf", ".txt"}
SAFE_INLINE_MEDIA_TYPES = {"application/pdf", "text/markdown", "text/plain"}
HTML_PREVIEW_EXTENSIONS = {".xlsx", ".xls", ".csv", ".docx"}
HTML_PREVIEW_CSP = (
    "default-src 'none'; "
    "style-src 'unsafe-inline'; "
    "base-uri 'none'; "
    "form-action 'none'; "
    "frame-ancestors 'self'"
)


def content_disposition_type_for_document(filename: str, media_type: str) -> str:
    ext = os.path.splitext(str(filename or ""))[1].lower()
    if ext in SAFE_INLINE_FILE_EXTENSIONS and media_type in SAFE_INLINE_MEDIA_TYPES:
        return "inline"
    return "attachment"


class KnowledgeDocumentContentService:
    """문서 원본/미리보기 응답 생성 경계.

    Endpoint는 권한과 resource resolve만 담당하고, 파일 fetch, HTML 변환,
    fallback response 생성은 이 service에서 관리한다.
    """

    def build_content_response(self, doc: Document) -> Response:
        if doc.source_type == "API" or not doc.file_path:
            raise HTTPException(
                status_code=400,
                detail="API로 받은 응답은 원문 보기를 제공하지 않습니다.",
            )

        file_path = str(doc.file_path)
        filename = str(doc.filename or "")
        is_external_file = self._is_external_file(file_path)
        if not is_external_file and not os.path.exists(file_path):
            raise HTTPException(status_code=404, detail="File not found on server")

        media_type, _ = mimetypes.guess_type(file_path)
        if not media_type:
            media_type = "application/octet-stream"
        content_disposition_type = content_disposition_type_for_document(
            filename,
            media_type,
        )

        ext = os.path.splitext(filename)[1].lower()
        if ext in HTML_PREVIEW_EXTENSIONS:
            preview_response = self._try_html_preview_response(
                file_path=file_path,
                ext=ext,
                is_external_file=is_external_file,
            )
            if preview_response is not None:
                return preview_response

        if is_external_file:
            return self._external_file_response(
                file_path=file_path,
                filename=filename,
                media_type=media_type,
                content_disposition_type=content_disposition_type,
            )

        return FileResponse(
            file_path,
            filename=filename,
            media_type=media_type,
            content_disposition_type=content_disposition_type,
            headers={"X-Content-Type-Options": "nosniff"},
        )

    def _try_html_preview_response(
        self,
        *,
        file_path: str,
        ext: str,
        is_external_file: bool,
    ) -> HTMLResponse | None:
        temp_file_path: str | None = None
        try:
            target_path = file_path
            if is_external_file:
                target_path, temp_file_path = self._download_external_preview_file(
                    file_path,
                    ext,
                )

            if ext == ".docx":
                body_content = self._docx_body_html(target_path)
            else:
                body_content = self._spreadsheet_body_html(target_path, ext)

            return self._html_preview_response(body_content)
        except EgressGuardError:
            raise
        except Exception as exc:
            logger.error("Document preview conversion failed: %s", type(exc).__name__)
            return None
        finally:
            if temp_file_path and os.path.exists(temp_file_path):
                try:
                    os.remove(temp_file_path)
                except Exception:
                    pass

    def _download_external_preview_file(self, file_path: str, ext: str) -> tuple[str, str | None]:
        if file_path.startswith("s3://"):
            # s3:// 직접 처리기는 아직 없다. 기존 fallback 경로에서 safe error로 닫는다.
            return file_path, None

        response = safe_http_request(
            "GET",
            file_path,
            operation_id=KNOWLEDGE_DOCUMENT_FETCH,
        )
        if response.status_code >= 400:
            raise RuntimeError("Remote file returned an error.")

        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            tmp.write(response.content)
            return tmp.name, tmp.name

    def _docx_body_html(self, file_path: str) -> str:
        doc_word = DocxDocument(file_path)
        paragraphs = [
            f"<p>{html.escape(paragraph.text)}</p>"
            for paragraph in doc_word.paragraphs
            if paragraph.text.strip()
        ]

        for table in doc_word.tables:
            rows_html: list[str] = []
            for row in table.rows:
                cells = [
                    f"<td>{html.escape(cell.text)}</td>"
                    for cell in row.cells
                ]
                rows_html.append(f"<tr>{''.join(cells)}</tr>")
            if rows_html:
                paragraphs.append(
                    f"<table class='docx-table'>{''.join(rows_html)}</table>"
                )
        return "\n".join(paragraphs)

    def _spreadsheet_body_html(self, file_path: str, ext: str) -> str:
        if ext == ".csv":
            df = pd.read_csv(file_path, nrows=100)
        else:
            df = pd.read_excel(file_path, nrows=100)
        return f"""
        <div class="info-banner">
            <span>⚠️</span>
            성능을 위해 상위 100행만 미리보기로 제공됩니다.
        </div>
        {df.to_html(index=False, border=0, escape=True)}
        """

    def _html_preview_response(self, body_content: str) -> HTMLResponse:
        html_content = f"""
        <!doctype html>
        <html>
        <head>
            <meta charset="utf-8" />
            <style>
                body {{ font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif; margin: 0; padding: 20px; background-color: #ffffff; line-height: 1.6; }}
                table {{ border-collapse: collapse; width: 100%; font-size: 14px; border: 1px solid #e5e7eb; margin-bottom: 20px; }}
                th {{ background-color: #f9fafb; color: #374151; font-weight: 600; text-align: left; padding: 12px 16px; border-bottom: 1px solid #e5e7eb; }}
                td {{ padding: 12px 16px; border-bottom: 1px solid #e5e7eb; color: #4b5563; }}
                tr:last-child td {{ border-bottom: none; }}
                tr:hover td {{ background-color: #f9fafb; }}
                p {{ margin-bottom: 0.8em; color: #1f2937; }}
                .docx-table td {{ border: 1px solid #e5e7eb; }}
                .info-banner {{
                    margin-bottom: 16px; padding: 10px 14px; background: #fffbeb; border: 1px solid #fcd34d;
                    color: #92400e; border-radius: 6px; font-size: 13px; font-weight: 500; display: flex; align-items: center; gap: 6px;
                }}
            </style>
        </head>
        <body>
            {body_content}
        </body>
        </html>
        """
        return HTMLResponse(
            content=html_content,
            headers={
                "Content-Security-Policy": HTML_PREVIEW_CSP,
                "X-Content-Type-Options": "nosniff",
            },
        )

    def _external_file_response(
        self,
        *,
        file_path: str,
        filename: str,
        media_type: str,
        content_disposition_type: str,
    ) -> StreamingResponse:
        try:
            external_res = safe_http_request(
                "GET",
                file_path,
                operation_id=KNOWLEDGE_DOCUMENT_FETCH,
            )
            if external_res.status_code >= 400:
                raise RuntimeError("Remote file returned an error.")

            def iterfile():
                yield external_res.content

            return StreamingResponse(
                iterfile(),
                media_type=media_type,
                headers={
                    "Content-Disposition": (
                        f"{content_disposition_type}; "
                        f"filename={safe_quote_filename(filename)}"
                    ),
                    "X-Content-Type-Options": "nosniff",
                },
            )
        except EgressGuardError as exc:
            logger.warning("External file proxy denied by egress guard: %s", exc.reason_code)
            raise HTTPException(
                status_code=400,
                detail={"reason_code": exc.reason_code},
            ) from exc
        except Exception as exc:
            logger.error("Failed to proxy external file: %s", type(exc).__name__)
            raise HTTPException(
                status_code=502,
                detail={"reason_code": "egress.proxy_failed"},
            ) from exc

    def _is_external_file(self, file_path: str) -> bool:
        return file_path.startswith("http") or file_path.startswith("s3://")
